import calendar
import io
import zipfile
from datetime import date

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .decorators import role_required
from .models import (
    Attendance, Event, EventMemberAttendance, EventStudentAttendance,
    Member, Session, Stream, Student, StudentAttendance,
)
from .views import can_manage_event

XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
HEADER_FILL = PatternFill('solid', fgColor='6C3FC5')
HEADER_FONT = Font(bold=True, color='FFFFFF')
ATT_HEADERS = ['Name', 'Roll Number', 'Branch', 'Present', 'Total Sessions', 'Percentage']


# ---------- small helpers ----------

def parse_date(value):
    """Turns '2026-09-01' into a date, or None if blank/invalid (never crashes)."""
    try:
        return date.fromisoformat(value) if value else None
    except ValueError:
        return None


def parse_period(request):
    """Returns (start, end, label). A ?month=YYYY-MM wins over start_date/end_date."""
    month = request.GET.get('month')
    if month:
        try:
            year, mon = (int(part) for part in month.split('-'))
            last_day = calendar.monthrange(year, mon)[1]
            return date(year, mon, 1), date(year, mon, last_day), f"{calendar.month_name[mon]}_{year}"
        except ValueError:
            pass
    start = parse_date(request.GET.get('start_date'))
    end = parse_date(request.GET.get('end_date'))
    if start or end:
        return start, end, f"{start or 'start'}_to_{end or 'today'}"
    return None, None, 'All_Time'


def object_from_query(request, model, key):
    """Finds e.g. the Stream whose id is in ?stream=3. Returns None if missing or invalid."""
    try:
        return model.objects.filter(id=int(request.GET.get(key, ''))).first()
    except ValueError:
        return None


def missing_choice(request, what):
    messages.error(request, f"Please choose a {what} first.")
    return redirect('report_downloads')


def percentage(present, total):
    return round((present / total) * 100, 1) if total else 0


def safe(value):
    """Stops text starting with = + - @ from being run as an Excel formula."""
    if isinstance(value, str) and value[:1] in ('=', '+', '-', '@'):
        return "'" + value
    return value


def add_sheet(wb, title, headers, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in rows:
        ws.append([safe(v) for v in row])
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 40)
    return ws


def workbook_bytes(wb):
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def xlsx_response(wb, filename):
    response = HttpResponse(workbook_bytes(wb), content_type=XLSX_TYPE)
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# ---------- report builders (data -> workbook) ----------

def stream_workbook(stream, start, end):
    """One workbook per stream: a Students sheet and a Committee Members sheet."""
    sessions = Session.objects.filter(stream=stream)
    if start:
        sessions = sessions.filter(session_date__gte=start)
    if end:
        sessions = sessions.filter(session_date__lte=end)
    total = sessions.count()

    student_rows = []
    for s in stream.students.all().order_by('full_name'):
        present = StudentAttendance.objects.filter(student=s, session__in=sessions, status='present').count()
        student_rows.append([s.full_name, s.roll_number, s.branch, present, total, percentage(present, total)])

    member_rows = []
    for m in stream.members.all().order_by('full_name'):
        present = Attendance.objects.filter(member=m, session__in=sessions, status='present').count()
        member_rows.append([m.full_name, m.roll_number, m.branch, present, total, percentage(present, total)])

    wb = Workbook()
    add_sheet(wb, 'Students', ATT_HEADERS, student_rows, first=True)
    add_sheet(wb, 'Committee Members', ATT_HEADERS, member_rows)
    return wb


def person_workbook(records, start, end):
    """
    One person's history. `records` is a StudentAttendance or Attendance queryset
    already filtered to that one person (both have the same fields).
    """
    records = records.select_related('session', 'session__stream').order_by('session__session_date')
    if start:
        records = records.filter(session__session_date__gte=start)
    if end:
        records = records.filter(session__session_date__lte=end)
    records = list(records)

    rows = [[r.session.session_date, r.session.stream.name, r.get_status_display()] for r in records]
    present = sum(1 for r in records if r.status == 'present')

    wb = Workbook()
    add_sheet(wb, 'Attendance', ['Date', 'Stream', 'Status'], rows, first=True)
    add_sheet(wb, 'Summary', ['Present', 'Total Sessions', 'Percentage'],
              [[present, len(records), percentage(present, len(records))]])
    return wb


def event_workbook(event):
    headers = ['Name', 'Roll Number', 'Status', 'Marked By', 'Marked At']

    def rows_for(queryset, person):
        return [
            [
                getattr(r, person).full_name,
                getattr(r, person).roll_number,
                r.get_status_display(),
                r.marked_by.username if r.marked_by else '',
                timezone.localtime(r.marked_at).strftime('%Y-%m-%d %H:%M'),
            ]
            for r in queryset
        ]

    students = (EventStudentAttendance.objects.filter(event=event)
                .select_related('student', 'marked_by').order_by('student__full_name'))
    members = (EventMemberAttendance.objects.filter(event=event)
               .select_related('member', 'marked_by').order_by('member__full_name'))

    wb = Workbook()
    add_sheet(wb, 'Students', headers, rows_for(students, 'student'), first=True)
    add_sheet(wb, 'Committee Members', headers, rows_for(members, 'member'))
    return wb


# ---------- views ----------

@role_required('admin', 'dairy')
def report_downloads(request):
    return render(request, 'core/report_downloads.html', {
        'streams': Stream.objects.all(),
        'students': Student.objects.order_by('full_name'),
        'members': Member.objects.order_by('full_name'),
        'events': Event.objects.order_by('-event_date'),
    })


@role_required('admin', 'dairy')
def download_stream_report(request):
    stream = object_from_query(request, Stream, 'stream')
    if stream is None:
        return missing_choice(request, 'stream')
    start, end, label = parse_period(request)
    return xlsx_response(stream_workbook(stream, start, end),
                         f"Attendance_{slugify(stream.name)}_{label}.xlsx")


@role_required('admin', 'dairy')
def download_all_reports(request):
    """Every stream's workbook, packed into one ZIP."""
    start, end, label = parse_period(request)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
        for stream in Stream.objects.all():
            zf.writestr(f"Attendance_{slugify(stream.name)}_{label}.xlsx",
                        workbook_bytes(stream_workbook(stream, start, end)))
    response = HttpResponse(buffer.getvalue(), content_type='application/zip')
    response['Content-Disposition'] = f'attachment; filename="Attendance_All_Streams_{label}.zip"'
    return response


@role_required('admin', 'dairy')
def download_student_report(request):
    student = object_from_query(request, Student, 'student')
    if student is None:
        return missing_choice(request, 'student')
    start, end, label = parse_period(request)
    wb = person_workbook(StudentAttendance.objects.filter(student=student), start, end)
    return xlsx_response(wb, f"Student_{slugify(student.roll_number)}_Attendance_{label}.xlsx")


@role_required('admin', 'dairy')
def download_member_report(request):
    member = object_from_query(request, Member, 'member')
    if member is None:
        return missing_choice(request, 'member')
    start, end, label = parse_period(request)
    wb = person_workbook(Attendance.objects.filter(member=member), start, end)
    return xlsx_response(wb, f"Member_{slugify(member.roll_number)}_Attendance_{label}.xlsx")


@login_required
def download_event_report(request):
    """Admin and Dairy: any event. Organizers: only events they can manage."""
    event = object_from_query(request, Event, 'event')
    if event is None:
        return missing_choice(request, 'event')
    if request.user.role != 'dairy' and not can_manage_event(request.user, event):
        raise PermissionDenied("You are not authorized to download this event's report.")
    return xlsx_response(event_workbook(event),
                         f"Event_Attendance_{slugify(event.name)}_{event.event_date}.xlsx")


@role_required('admin', 'dairy')
def download_all_event_reports(request):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
        for event in Event.objects.all():
            zf.writestr(f"Event_Attendance_{slugify(event.name)}_{event.event_date}_{event.id}.xlsx",
                        workbook_bytes(event_workbook(event)))
    response = HttpResponse(buffer.getvalue(), content_type='application/zip')
    response['Content-Disposition'] = 'attachment; filename="Event_Attendance_All_Events.zip"'
    return response